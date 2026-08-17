"""Build a .pptx deck from the structured slide JSON with python-pptx, emitting every
layout as native editable shapes rather than a reference deck.
"""

import math
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from slides.icons import ICON_FONT, add_icon_run, resolve_icon, section_glyph

USE_EMOJI = False

# ---- Canvas (16:9) ----
SLIDE_W = 12192000
SLIDE_H = 6858000
MARGIN = 610000

NAV_TOP = 246000
NAV_H = 300000
NAV_LABEL_H = 250000
TITLE_TOP = 730000
TITLE_SIZE = 28
BODY_TOP = 1880000
BODY_BOT = 430000
BODY_H = SLIDE_H - BODY_TOP - BODY_BOT

RIGHT_X = 7560000
RIGHT_W = SLIDE_W - RIGHT_X - MARGIN
LEFT_W = RIGHT_X - MARGIN - 240000

FONT = "Segoe UI"
EMOJI_FONT = "Segoe UI Emoji"

TITLE_PT = 32
WORD_LABEL_PT = 16
# Fraction of white mixed into a colour for a panel fill: 0.88 leaves 12% colour.
PANEL_TINT = 0.88

# ---- Palette (broad Office-style set) ----
INK = RGBColor(0x26, 0x26, 0x26)
MUTED = RGBColor(0x7F, 0x7F, 0x7F)
HAIRLINE = RGBColor(0xD9, 0xD9, 0xD9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0xF2, 0xF4, 0xF7)
PRIMARY = RGBColor(0x2E, 0x75, 0xB6)

CATEGORICAL = [
    RGBColor(0x2E, 0x75, 0xB6),  # blue
    RGBColor(0xED, 0x7D, 0x31),  # orange
    RGBColor(0x70, 0xAD, 0x47),  # green
    RGBColor(0xFF, 0xC0, 0x00),  # gold
    RGBColor(0x70, 0x30, 0xA0),  # purple
    RGBColor(0xC0, 0x00, 0x00),  # red
]


def build_pptx(data: dict, output_path: Path, config: dict) -> None:
    use_emoji = str(config.get("render", {}).get("icons", "fluent")).lower() == "emoji"

    meta = data.get("meta", {})
    sections = data.get("sections", [])
    slides = data.get("slides", [])
    defaults = config.get("defaults", {})

    author = meta.get("author") or defaults.get("author", "")
    affiliation = meta.get("affiliation") or defaults.get("affiliation", "")
    email = meta.get("email") or defaults.get("email", "")
    venue = meta.get("venue", "") or defaults.get("venue", "")
    title = meta.get("title", "Untitled Presentation")

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    blank = prs.slide_layouts[6]

    _title_slide(prs.slides.add_slide(blank), title, author, affiliation, venue)
    total = len(slides)
    for i, slide_data in enumerate(slides, 1):
        print(f"       [{i}/{total}] {slide_data.get('title', 'Slide')}", file=sys.stderr)
        slide = prs.slides.add_slide(blank)
        slide_title = slide_data.get("title", "Slide")
        section = slide_data.get("section", sections[0] if sections else "")
        layout = slide_data.get("layout", "bullets")
        body = slide_data.get("body", {}) or {}
        _nav(slide, sections, section)
        _title(slide, slide_title, size=TITLE_SIZE, icon_glyph=section_glyph(section))
        _draw_body(slide, layout, body, use_emoji)
    _closing_slide(prs.slides.add_slide(blank), author, email, venue)

    prs.save(str(output_path))


# ---- low-level helpers ----

def _area(left, top, width, height) -> dict:
    """One rectangle of the slide, in EMU."""
    return {"left": left, "top": top, "width": width, "height": height}


def _box(slide, left, top, width, height) -> dict:
    """A borderless textbox, as the `shape` itself and its text `frame`."""
    tb = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    return {"shape": tb, "frame": tf}


def _shape(slide, kind, box):
    """One shape of the given kind, with PowerPoint's default shadow turned off."""
    sp = slide.shapes.add_shape(kind, Emu(box["left"]), Emu(box["top"]),
                                Emu(box["width"]), Emu(box["height"]))
    sp.shadow.inherit = False
    return sp


def _rect(slide, left, top, width, height):
    """A plain rectangle."""
    return _shape(slide, MSO_SHAPE.RECTANGLE, _area(left, top, width, height))


def _rounded(slide, left, top, width, height):
    """A rounded rectangle, the shape every panel and card uses."""
    return _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, _area(left, top, width, height))


def _oval(slide, left, top, width, height):
    """An ellipse, used for the icon discs."""
    return _shape(slide, MSO_SHAPE.OVAL, _area(left, top, width, height))


def _solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _no_fill(shape):
    shape.fill.background()
    shape.line.fill.background()


def _tint(color, amount=PANEL_TINT):
    """Lighten a colour towards white (amount = fraction of white)."""
    r = int(color[0] + (255 - color[0]) * amount)
    g = int(color[1] + (255 - color[1]) * amount)
    b = int(color[2] + (255 - color[2]) * amount)
    return RGBColor(r, g, b)


def _para(tf, first=False):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    return p


def _add_runs(p, runs):
    """runs: list of (text, bold, color)."""
    for text, bold, color in runs:
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.bold = bool(bold)
        r.font.color.rgb = color or INK


def _md_runs(text, base=INK, accent=PRIMARY):
    """Parse **bold** markdown into runs; bold terms get the accent colour."""
    out = []
    for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", str(text))):
        if not seg:
            continue
        if i % 2 == 1:
            out.append((seg, True, accent))
        else:
            out.append((seg, False, base))
    return out or [(str(text), False, base)]


def _set_bullet(p, color=PRIMARY, char="●"):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "320040")
    pPr.set("indent", "-320040")
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        e = pPr.find(qn(tag))
        if e is not None:
            pPr.remove(e)
    buClr = pPr.makeelement(qn("a:buClr"), {})
    srgb = pPr.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % (color[0], color[1], color[2])})
    buClr.append(srgb)
    pPr.append(buClr)
    pPr.append(pPr.makeelement(qn("a:buSzPct"), {"val": "70000"}))
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": char}))


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        e = pPr.find(qn(tag))
        if e is not None:
            pPr.remove(e)
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


# ---- chrome: nav breadcrumb, title, rule ----

def _nav(slide, sections, current):
    n = max(1, len(sections))
    nav_w = SLIDE_W - 2 * MARGIN
    seg_w = nav_w // n
    for i, s in enumerate(sections):
        active = (s == current)
        seg_l = MARGIN + i * seg_w
        tf = _box(slide, seg_l, NAV_TOP, seg_w, NAV_LABEL_H)["frame"]
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = s
        r.font.name = FONT
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = INK if active else MUTED
        # underline track
        line = _rect(slide, seg_l + 20000, NAV_TOP + NAV_LABEL_H, seg_w - 40000, 48000)
        _solid(line, PRIMARY if active else HAIRLINE)


def _title(slide, text, size=TITLE_PT, icon_glyph=None):
    tf = _box(slide, MARGIN, TITLE_TOP, SLIDE_W - 2 * MARGIN, BODY_TOP - TITLE_TOP)["frame"]
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    _no_bullet(p)
    if icon_glyph:
        ir = p.add_run(); ir.text = icon_glyph + "  "
        ir.font.name = ICON_FONT; ir.font.size = Pt(size - 4); ir.font.color.rgb = PRIMARY
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.bold = True
    r.font.size = Pt(size)
    r.font.color.rgb = INK
    width_emu = SLIDE_W - 2 * MARGIN
    char_w = max(1, size * 0.52 * 12700)
    cpl = max(12, int(width_emu / char_w))
    eff_len = len(text) + (3 if icon_glyph else 0)
    lines = max(1, math.ceil(eff_len / cpl))
    line_h = int(size * 1.18 * 12700)
    rule_y = min(TITLE_TOP + lines * line_h + 80000, BODY_TOP - 150000)
    _solid(_rect(slide, MARGIN, rule_y, 880000, 52000), PRIMARY)


# ---- body font sizing (deterministic shrink-to-fit) ----

def _body_size(words):
    if words <= 24:
        return 24
    if words <= 38:
        return 21
    if words <= 52:
        return 19
    if words <= 66:
        return 17
    return 16


def _bullets(slide, bullets, box, size=None, anchor=MSO_ANCHOR.TOP):
    """Draw the bullet list inside `box`, sized to fit unless `size` says otherwise."""
    words = sum(len(str(b.get("text", "")).split()) for b in bullets)
    if size is None:
        size = _body_size(words)
    spc = 16 if len(bullets) <= 3 and words <= 30 else 11
    tf = _box(slide, box["left"], box["top"], box["width"], box["height"])["frame"]
    tf.vertical_anchor = anchor
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    first = True
    for b in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(spc)
        p.line_spacing = 1.15
        _set_bullet(p, PRIMARY)
        _add_runs(p, _md_runs(b.get("text", "")))
        for r in p.runs:
            r.font.size = Pt(size)


# ---- evidence: stat / quote / note ----

def _stat(slide, number, label, at, color):
    """One big number over its label, placed at the `left`/`top` corner in `at`."""
    tf = _box(slide, at["left"], at["top"], RIGHT_W, 1000000)["frame"]
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _no_bullet(p)
    r = p.add_run(); r.text = number
    r.font.name = FONT; r.font.bold = True; r.font.size = Pt(40); r.font.color.rgb = color
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    _no_bullet(p2)
    r2 = p2.add_run(); r2.text = label
    r2.font.name = FONT; r2.font.bold = True; r2.font.size = Pt(13); r2.font.color.rgb = INK


def _evidence_box(slide, box, body_runs_list, style):
    """A tinted panel holding a quote or a note.

    Args:
        slide: the slide to draw on.
        box: the area to fill, as `left`, `top`, `width` and `height`.
        body_runs_list: one list of styled runs per paragraph of the body.
        style: `color`, and optionally `label`, `source`, `italic` and `icon`.
    """
    color = style["color"]
    label = style.get("label")
    source = style.get("source")
    italic = style.get("italic", True)
    icon = style.get("icon")
    sp = _rounded(slide, box["left"], box["top"], box["width"], box["height"])
    sp.fill.solid(); sp.fill.fore_color.rgb = _tint(color, 0.93)
    sp.line.color.rgb = color; sp.line.width = Pt(1.5)
    tf = sp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Emu(180000)
    tf.margin_top = tf.margin_bottom = Emu(150000)
    first = True
    if label:
        p = tf.paragraphs[0]; first = False
        _no_bullet(p)
        ic = add_icon_run(p, icon, 16, color)
        if ic is not None:
            p.add_run().text = "  "
        r = p.add_run(); r.text = label
        r.font.name = FONT; r.font.bold = True; r.font.size = Pt(16); r.font.color.rgb = color
    elif icon:
        p = tf.paragraphs[0]; first = False
        _no_bullet(p)
        add_icon_run(p, icon, 22, color)
    for runs in body_runs_list:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        _no_bullet(p); p.line_spacing = 1.12
        for text, bold, color2 in runs:
            r = p.add_run(); r.text = text
            r.font.name = FONT; r.font.bold = bold; r.font.italic = italic
            r.font.size = Pt(19); r.font.color.rgb = color2 or INK
    if source:
        p = tf.add_paragraph(); p.alignment = PP_ALIGN.RIGHT
        _no_bullet(p); p.space_before = Pt(6)
        r = p.add_run(); r.text = source
        r.font.name = FONT; r.font.size = Pt(13); r.font.color.rgb = MUTED


def _stat_trio(slide, items, box):
    """Two or three big numbers stacked vertically in the right column.

    Args:
        slide: the slide to draw on.
        items: up to three mappings with `number` and `label`.
        box: the area to fill, as `left`, `top`, `width` and `height`.
    """
    items = items[:3]
    n = max(1, len(items))
    gap = 200000
    cell_h = (box["height"] - (n - 1) * gap) // n
    for i, it in enumerate(items):
        cy = box["top"] + i * (cell_h + gap)
        color = CATEGORICAL[i % len(CATEGORICAL)]
        tf = _box(slide, box["left"], cy, box["width"], cell_h)["frame"]
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _no_bullet(p)
        r = p.add_run(); r.text = str(it.get("number", ""))
        r.font.name = FONT; r.font.bold = True; r.font.size = Pt(40); r.font.color.rgb = color
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; _no_bullet(p2)
        r2 = p2.add_run(); r2.text = str(it.get("label", ""))
        r2.font.name = FONT; r2.font.bold = True; r2.font.size = Pt(13); r2.font.color.rgb = INK


def _render_evidence(slide, ev):
    """Right-column evidence: stats / stat / quote / note / figure."""
    if not ev:
        return
    kind = ev.get("kind", "note")
    if kind == "stats":
        _stat_trio(slide, ev.get("items", []), {
            "left": RIGHT_X,
            "top": BODY_TOP,
            "width": RIGHT_W,
            "height": BODY_H,
        })
        return
    if kind == "stat":
        _stat(slide, ev.get("number", ev.get("text", "")), ev.get("label", ev.get("source", "")),
              {"left": RIGHT_X, "top": BODY_TOP + 200000}, CATEGORICAL[0])
        return
    if kind == "quote-green" or kind == "quote-dark":
        txt = ev.get("text", "")
        _evidence_box(slide, _area(RIGHT_X, BODY_TOP, RIGHT_W, BODY_H),
                      [[("“" + txt + "”", False, INK)]],
                      {
                          "color": PRIMARY,
                          "source": ("· " + ev["source"]) if ev.get("source") else None,
                          "italic": True,
                          "icon": "comment",
                      })
        return
    if kind == "note":
        _evidence_box(slide, _area(RIGHT_X, BODY_TOP, RIGHT_W, BODY_H),
                      [_md_runs(ev.get("text", ""))],
                      {
                          "color": CATEGORICAL[1],
                          "label": ev.get("label") or "Note",
                          "source": ev.get("source"),
                          "italic": False,
                          "icon": "idea",
                      })
        return
    # figure placeholder
    sp = _rounded(slide, RIGHT_X, BODY_TOP, RIGHT_W, BODY_H - 200000)
    sp.fill.solid(); sp.fill.fore_color.rgb = PANEL; sp.line.color.rgb = HAIRLINE
    tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _no_bullet(p)
    r = p.add_run(); r.text = "[ figure: " + (ev.get("alt") or ev.get("text") or "") + " ]"
    r.font.name = FONT; r.font.size = Pt(13); r.font.color.rgb = MUTED


# ---- layouts ----

def _layout_split(slide, body):
    bullets = body.get("bullets", [])
    ev = body.get("evidence") or {}
    if ev:
        _bullets(slide, bullets, _area(MARGIN, BODY_TOP, LEFT_W, BODY_H))
        _render_evidence(slide, ev)
    else:
        _bullets(slide, bullets, _area(MARGIN, BODY_TOP, SLIDE_W - 2 * MARGIN, BODY_H))


def _layout_bullets(slide, body):
    _layout_split(slide, body)


def _layout_iconrow(slide, body, use_emoji=False):
    cols = body.get("columns", [])[:3]
    n = max(1, len(cols))
    gap = 280000
    col_w = (SLIDE_W - 2 * MARGIN - (n - 1) * gap) // n
    disc = 1000000
    for i, col in enumerate(cols):
        cx = MARGIN + i * (col_w + gap)
        color = CATEGORICAL[i % len(CATEGORICAL)]
        icon = resolve_icon(col.get("icon"), use_emoji)
        ch, font = icon["char"], icon["font"]

        circ = _oval(slide, int(cx + col_w / 2 - disc / 2), BODY_TOP, disc, disc)
        circ.fill.solid(); circ.fill.fore_color.rgb = _tint(color, 0.86)
        circ.line.color.rgb = color; circ.line.width = Pt(1)
        ctf = circ.text_frame; ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        ep = ctf.paragraphs[0]; ep.alignment = PP_ALIGN.CENTER; _no_bullet(ep)
        if ch:
            er = ep.add_run(); er.text = ch
            er.font.name = font; er.font.size = Pt(40)
            if font == ICON_FONT:
                er.font.color.rgb = color
        # text
        tf = _box(slide, cx, BODY_TOP + disc + 120000, col_w, BODY_H - disc - 120000)["frame"]
        lead, text, quote = col.get("lead", ""), col.get("text", ""), col.get("quote", "")
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _no_bullet(p)
        runs = []
        if lead:
            runs.append((lead + ("  " if text else ""), True, INK))
        if text:
            runs.append((text, False, INK))
        _add_runs(p, runs or [("", False, INK)])
        for r in p.runs:
            r.font.size = Pt(15)
        if quote:
            pq = tf.add_paragraph(); pq.alignment = PP_ALIGN.CENTER; _no_bullet(pq)
            rq = pq.add_run(); rq.text = quote
            rq.font.name = FONT; rq.font.italic = True; rq.font.size = Pt(13); rq.font.color.rgb = MUTED


def _two_col(slide, left_paras, right_paras):
    gap = 360000
    col_w = (SLIDE_W - 2 * MARGIN - gap) // 2
    for idx, paras in enumerate((left_paras, right_paras)):
        cx = MARGIN + idx * (col_w + gap)
        size = _body_size(sum(len(str(b.get("text", "")).split()) for b in paras))
        tf = _box(slide, cx, BODY_TOP, col_w, BODY_H)["frame"]
        first = True
        for j, b in enumerate(paras):
            p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
            p.space_after = Pt(8); p.line_spacing = 1.08
            heading = b.get("heading")
            if j == 0 and heading:
                _no_bullet(p)
                r = p.add_run(); r.text = heading
                r.font.name = FONT; r.font.bold = True; r.font.size = Pt(18); r.font.color.rgb = CATEGORICAL[idx % len(CATEGORICAL)]
                continue
            _set_bullet(p, CATEGORICAL[idx % len(CATEGORICAL)])
            _add_runs(p, _md_runs(b.get("text", "")))
            for r in p.runs:
                r.font.size = Pt(size)
    # divider
    dl = MARGIN + col_w + gap // 2
    line = _rect(slide, dl, BODY_TOP + 60000, 9525, BODY_H - 120000)
    _solid(line, HAIRLINE)


def _layout_cards(slide, body):
    cards = body.get("cards", [])[:2]
    n = max(1, len(cards))
    gap = 360000
    col_w = (SLIDE_W - 2 * MARGIN - (n - 1) * gap) // n
    for i, card in enumerate(cards):
        cx = MARGIN + i * (col_w + gap)
        color = CATEGORICAL[i % len(CATEGORICAL)]
        sp = _rounded(slide, cx, BODY_TOP, col_w, BODY_H)
        sp.fill.solid(); sp.fill.fore_color.rgb = _tint(color, 0.92)
        sp.line.color.rgb = color; sp.line.width = Pt(1.25)
        # header bar
        hdr = _rounded(slide, cx, BODY_TOP, col_w, 560000)
        _solid(hdr, color)
        htf = hdr.text_frame; htf.vertical_anchor = MSO_ANCHOR.MIDDLE
        htf.margin_left = Emu(200000)
        hp = htf.paragraphs[0]; _no_bullet(hp)
        ic = add_icon_run(hp, card.get("icon"), 18, WHITE)
        if ic is not None:
            hp.add_run().text = "  "
        hr = hp.add_run(); hr.text = card.get("title", "")
        hr.font.name = FONT; hr.font.bold = True; hr.font.size = Pt(18); hr.font.color.rgb = WHITE
        # body
        tf = _box(slide, cx + 200000, BODY_TOP + 700000, col_w - 400000, BODY_H - 880000)["frame"]
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; _no_bullet(p); p.line_spacing = 1.12
        _add_runs(p, _md_runs(card.get("text", "")))
        for r in p.runs:
            r.font.size = Pt(16)


def _layout_matrix(slide, body):
    cells = body.get("cells", [])[:4]
    gap = 280000
    cw = (SLIDE_W - 2 * MARGIN - gap) // 2
    ch = (BODY_H - gap) // 2
    for i, cell in enumerate(cells):
        r_, c_ = divmod(i, 2)
        cx = MARGIN + c_ * (cw + gap)
        cy = BODY_TOP + r_ * (ch + gap)
        color = CATEGORICAL[i % len(CATEGORICAL)]
        sp = _rounded(slide, cx, cy, cw, ch)
        sp.fill.solid(); sp.fill.fore_color.rgb = _tint(color, 0.9)
        sp.line.color.rgb = color; sp.line.width = Pt(6)
        tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = Emu(180000); tf.margin_top = Emu(160000)
        p = tf.paragraphs[0]; _no_bullet(p)
        ic = add_icon_run(p, cell.get("icon"), 17, color)
        if ic is not None:
            p.add_run().text = "  "
        r = p.add_run(); r.text = cell.get("label", "")
        r.font.name = FONT; r.font.bold = True; r.font.size = Pt(18); r.font.color.rgb = color
        if cell.get("text"):
            p2 = tf.add_paragraph(); _no_bullet(p2); p2.line_spacing = 1.1; p2.space_before = Pt(4)
            _add_runs(p2, _md_runs(cell["text"]))
            for rr in p2.runs:
                rr.font.size = Pt(15)


def _layout_stepflow(slide, body):
    steps = body.get("steps", [])
    n = max(1, len(steps))
    gap = 360000
    sw = (SLIDE_W - 2 * MARGIN - (n - 1) * gap) // n
    sh = min(BODY_H, 2200000)
    sy = BODY_TOP + (BODY_H - sh) // 2
    for i, step in enumerate(steps):
        sx = MARGIN + i * (sw + gap)
        color = CATEGORICAL[i % len(CATEGORICAL)]
        sp = _rounded(slide, sx, sy, sw, sh)
        sp.fill.solid(); sp.fill.fore_color.rgb = PANEL
        sp.line.color.rgb = color; sp.line.width = Pt(1.5)
        tf = sp.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Emu(140000)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _no_bullet(p)
        _add_runs(p, _md_runs(str(step)))
        for r in p.runs:
            r.font.size = Pt(15)
        if i < n - 1:
            atf = _box(slide, sx + sw, sy, gap, sh)["frame"]
            atf.vertical_anchor = MSO_ANCHOR.MIDDLE
            ap = atf.paragraphs[0]; ap.alignment = PP_ALIGN.CENTER; _no_bullet(ap)
            ar = ap.add_run(); ar.text = "→"
            ar.font.name = FONT; ar.font.size = Pt(28); ar.font.color.rgb = MUTED


def _layout_statement(slide, body):
    tf = _box(slide, MARGIN, BODY_TOP, SLIDE_W - 2 * MARGIN, BODY_H)["frame"]
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; _no_bullet(p)
    _add_runs(p, _md_runs(body.get("claim", "")))
    for r in p.runs:
        r.font.size = Pt(30); r.font.bold = True
    if body.get("sub"):
        p2 = tf.add_paragraph(); _no_bullet(p2); p2.space_before = Pt(14)
        _add_runs(p2, _md_runs(body["sub"]))
        for r in p2.runs:
            r.font.size = Pt(18); r.font.color.rgb = MUTED


def _layout_boxes(slide, body):
    _bullets(slide, body.get("bullets", []), _area(MARGIN, BODY_TOP, SLIDE_W - 2 * MARGIN,
                                              BODY_H))


def _word_label(slide, text, centre, color, size=WORD_LABEL_PT):
    """One word centred on the `cx`/`cy` point in `centre`."""
    w, h = 1900000, 380000
    tf = _box(slide, int(centre["cx"] - w / 2), int(centre["cy"] - h / 2), w, h)["frame"]
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _no_bullet(p)
    r = p.add_run(); r.text = text
    r.font.name = FONT; r.font.bold = True; r.font.size = Pt(size); r.font.color.rgb = color


def _layout_circle(slide, body):
    """Ring with a centre word and orbiting keywords; optional bullets on the left."""
    bullets = body.get("bullets", [])
    center = body.get("center", "")
    keywords = body.get("keywords", [])
    if bullets:
        _bullets(slide, bullets, _area(MARGIN, BODY_TOP, LEFT_W, BODY_H))
        ring_cx = RIGHT_X + RIGHT_W // 2
        ring_d = min(RIGHT_W, BODY_H) - 200000
    else:
        ring_cx = SLIDE_W // 2
        ring_d = min(BODY_H, 3400000)
    ring_cy = BODY_TOP + BODY_H // 2
    oval = _oval(slide, int(ring_cx - ring_d / 2), int(ring_cy - ring_d / 2),
                 ring_d, ring_d)
    oval.fill.background()
    oval.line.color.rgb = PRIMARY
    oval.line.width = Pt(2)
    if center:
        _word_label(slide, center, {"cx": ring_cx, "cy": ring_cy}, INK, size=20)
    r = ring_d / 2 + 120000
    n = max(1, len(keywords))
    for i, kw in enumerate(keywords):
        ang = -math.pi / 2 + i * (2 * math.pi / n)
        kx = ring_cx + r * math.cos(ang)
        ky = ring_cy + r * math.sin(ang)
        _word_label(slide, kw, {"cx": kx, "cy": ky},
                    CATEGORICAL[i % len(CATEGORICAL)], size=15)


def _draw_body(slide, layout, body, use_emoji):
    """Draw one slide's body with the layout it asked for.

    Only the icon row cares whether icons render as emoji, so it is the one
    layout handed that choice.
    """
    if layout == "iconrow":
        return _layout_iconrow(slide, body, use_emoji)
    return _LAYOUTS.get(layout, _layout_bullets)(slide, body)


_LAYOUTS = {
    "split": _layout_split,
    "bullets": _layout_bullets,
    "iconrow": _layout_iconrow,
    "cards": _layout_cards,
    "matrix": _layout_matrix,
    "stepflow": _layout_stepflow,
    "statement": _layout_statement,
    "boxes": _layout_boxes,
    "circle": _layout_circle,
    "prose": _layout_split,
}


# ---- title & closing ----

def _title_slide(slide, title, author, affiliation, venue):
    _rect_full = _rect(slide, 0, 0, SLIDE_W, 120000)
    _solid(_rect_full, PRIMARY)
    if venue:
        vtf = _box(slide, MARGIN, 300000, SLIDE_W - 2 * MARGIN, 400000)["frame"]
        vtf.vertical_anchor = MSO_ANCHOR.MIDDLE
        vp = vtf.paragraphs[0]; _no_bullet(vp)
        vr = vp.add_run(); vr.text = venue
        vr.font.name = FONT; vr.font.bold = True; vr.font.size = Pt(15); vr.font.color.rgb = MUTED
    # centred title block
    tf = _box(slide, MARGIN, 2150000, SLIDE_W - 2 * MARGIN, 2400000)["frame"]
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tp = tf.paragraphs[0]; tp.alignment = PP_ALIGN.CENTER; _no_bullet(tp)
    tr = tp.add_run(); tr.text = title
    tr.font.name = FONT; tr.font.bold = True; tr.font.size = Pt(38); tr.font.color.rgb = INK
    ap = tf.add_paragraph(); ap.alignment = PP_ALIGN.CENTER; _no_bullet(ap); ap.space_before = Pt(20)
    ar = ap.add_run(); ar.text = author
    ar.font.name = FONT; ar.font.italic = True; ar.font.size = Pt(20); ar.font.color.rgb = INK
    if affiliation:
        afp = tf.add_paragraph(); afp.alignment = PP_ALIGN.CENTER; _no_bullet(afp)
        afr = afp.add_run(); afr.text = affiliation
        afr.font.name = FONT; afr.font.size = Pt(15); afr.font.color.rgb = MUTED


def _closing_slide(slide, author, email, venue):
    rb = _rect(slide, 0, 0, SLIDE_W, 120000)
    _solid(rb, PRIMARY)
    tf = _box(slide, MARGIN, 2400000, SLIDE_W - 2 * MARGIN, 2000000)["frame"]
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; _no_bullet(p)
    r = p.add_run(); r.text = "Thank you"
    r.font.name = FONT; r.font.bold = True; r.font.size = Pt(48); r.font.color.rgb = INK
    contact = author + (("  ·  " + email) if email else "")
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; _no_bullet(p2); p2.space_before = Pt(18)
    r2 = p2.add_run(); r2.text = contact
    r2.font.name = FONT; r2.font.size = Pt(18); r2.font.color.rgb = MUTED
